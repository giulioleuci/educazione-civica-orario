import os
from collections import defaultdict
from datetime import datetime, timedelta

import pandas as pd
from openpyxl.styles import PatternFill, Border, Side, Alignment, Font
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.pagebreak import Break
from openpyxl.worksheet.properties import PageSetupProperties

from .sicurezza import _sanitize_for_excel, _sanitize_output_path, _sanitize_sheet_name


_COLORE_SCURO = '1F3864'

_COLORE_BORDO = 'BFBFBF'

_NOMI_GIORNI = {'LUN': 'Lunedì', 'MAR': 'Martedì', 'MER': 'Mercoledì', 'GIO': 'Giovedì', 'VEN': 'Venerdì', 'SAB': 'Sabato'}

def _get_excel_styles():
    """Restituisce gli stili predefiniti per i fogli Excel."""
    return {
        'header_fill': PatternFill(start_color='B8CCE4', end_color='B8CCE4', fill_type='solid'),
        'week_fill': PatternFill(start_color='E6E6FA', end_color='E6E6FA', fill_type='solid'),
        'thin_border': Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        ),
        'centered_alignment': Alignment(horizontal='center', vertical='center'),
        'bold_font': Font(bold=True)
    }

def _get_week_range(date):
    """Funzione per calcolare il range settimanale (lun-sab)"""
    start = date - timedelta(days=date.weekday())  # Lunedì
    end = start + timedelta(days=5)  # Sabato
    return f"{start.strftime('%d/%m/%Y')} - {end.strftime('%d/%m/%Y')}"

def _scrivi_titolo_foglio(ws, titolo, sottotitolo, n_colonne):
    """Scrive titolo (riga 1) e sottotitolo (riga 2) uniti su tutte le colonne del foglio"""
    for riga, testo, font, altezza in ((1, titolo, Font(bold=True, size=16, color=_COLORE_SCURO), 28),
                                       (2, sottotitolo, Font(italic=True, size=10, color='595959'), 18)):
        ws.merge_cells(start_row=riga, start_column=1, end_row=riga, end_column=n_colonne)
        ws.cell(row=riga, column=1, value=testo)
        ws.cell(row=riga, column=1).font = font
        ws.cell(row=riga, column=1).alignment = Alignment(horizontal='left', vertical='center')
        ws.row_dimensions[riga].height = altezza

def genera_orario_classi(calendario, classi_df, cartella_output):
    """Genera il file orario_classi.xlsx: un foglio per ogni classe, con le sostituzioni settimanali"""

    # Stili del foglio classe (tavolozza blu scuro / grigio, bordi sottili chiari, righe alternate)
    colore_scuro = _COLORE_SCURO
    colore_bordo = _COLORE_BORDO
    header_fill = PatternFill(start_color=colore_scuro, end_color=colore_scuro, fill_type='solid')
    banda_fill = PatternFill(start_color='EEF3FA', end_color='EEF3FA', fill_type='solid')
    lato = Side(style='thin', color=colore_bordo)
    bordo = Border(left=lato, right=lato, top=lato, bottom=lato)
    bordo_nuovo_mese = Border(left=lato, right=lato, bottom=lato, top=Side(style='medium', color=colore_scuro))
    centrato = Alignment(horizontal='center', vertical='center')
    a_sinistra = Alignment(horizontal='left', vertical='center', indent=1)
    font_header = Font(bold=True, color='FFFFFF')

    colonne = ['N.', 'Settimana', 'Data', 'Giorno', 'Ora', 'Docente di civica', 'Docente sostituito']
    colonne_a_sinistra = {6, 7}  # nomi dei docenti
    riga_header = 4

    writer_classi = pd.ExcelWriter(os.path.join(cartella_output, 'orario_classi.xlsx'), engine='openpyxl')

    # Pre-raggruppamento delle voci per classe per ottimizzare la ricerca
    calendario_per_classe = defaultdict(list)
    for entry in calendario:
        calendario_per_classe[entry['CLASSE']].append(entry)

    # Per ogni classe, estraiamo le entries corrispondenti e creiamo un foglio
    for nome_classe in classi_df['CLASSE']:
        class_entries = sorted(calendario_per_classe.get(nome_classe, []), key=lambda x: x['DATA'])

        class_data = []
        for numero, entry in enumerate(class_entries, start=1):
            class_data.append({
                'N.': numero,
                'Settimana': _get_week_range(entry['DATA']),
                'Data': entry['DATA'].date(),
                'Giorno': _NOMI_GIORNI.get(entry['GIORNO'], entry['GIORNO']),
                'Ora': entry['ORA'],
                'Docente di civica': entry['DOCENTE_CIVICS'],
                'Docente sostituito': entry['DOCENTE_SOSTITUITO']
            })

        if class_data:
            df = pd.DataFrame(class_data, columns=colonne)
            df = _sanitize_for_excel(df)
            safe_sheet_name = _sanitize_sheet_name(nome_classe, default="Classe")
            df.to_excel(writer_classi, sheet_name=safe_sheet_name, index=False, startrow=riga_header - 1)

            ws = writer_classi.sheets[safe_sheet_name]
            ultima_riga = riga_header + len(class_data)
            n_colonne = len(colonne)

            # Titolo e sottotitolo (docente/i, numero di ore, periodo coperto)
            docenti = sorted({e['DOCENTE_CIVICS'] for e in class_entries})
            _scrivi_titolo_foglio(
                ws, f"Educazione civica – Classe {nome_classe}",
                f"Docente di civica: {', '.join(docenti)}  ·  {len(class_entries)} ore  ·  "
                f"dal {class_entries[0]['DATA'].strftime('%d/%m/%Y')} al {class_entries[-1]['DATA'].strftime('%d/%m/%Y')}",
                n_colonne)
            ws.row_dimensions[riga_header].height = 24

            # Intestazione della tabella
            for col in range(1, n_colonne + 1):
                cell = ws.cell(row=riga_header, column=col)
                cell.fill = header_fill
                cell.font = font_header
                cell.border = bordo
                cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)

            # Righe dati: righe alternate, data in formato italiano, bordo marcato al cambio di mese
            mese_precedente = None
            for indice, entry in enumerate(class_entries):
                riga = riga_header + 1 + indice
                nuovo_mese = mese_precedente is not None and entry['DATA'].month != mese_precedente
                mese_precedente = entry['DATA'].month
                for col in range(1, n_colonne + 1):
                    cell = ws.cell(row=riga, column=col)
                    cell.border = bordo_nuovo_mese if nuovo_mese else bordo
                    cell.alignment = a_sinistra if col in colonne_a_sinistra else centrato
                    if indice % 2 == 1:
                        cell.fill = banda_fill
                ws.cell(row=riga, column=3).number_format = 'DD/MM/YYYY'

            # Larghezza colonne in base al contenuto della tabella (titolo escluso)
            for col in range(1, n_colonne + 1):
                lunghezza = max(len(str(ws.cell(row=riga, column=col).value or ''))
                                for riga in range(riga_header, ultima_riga + 1))
                ws.column_dimensions[get_column_letter(col)].width = max(lunghezza + 4, 8)

            # Vista e stampa: intestazione fissa, filtri, niente griglia, A4 in larghezza con intestazione ripetuta
            ws.freeze_panes = ws.cell(row=riga_header + 1, column=1)
            ws.auto_filter.ref = f"A{riga_header}:{get_column_letter(n_colonne)}{ultima_riga}"
            ws.sheet_view.showGridLines = False
            ws.sheet_properties.tabColor = colore_scuro
            ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
            ws.page_setup.orientation = 'portrait'
            ws.page_setup.paperSize = ws.PAPERSIZE_A4
            ws.page_setup.fitToWidth = 1
            ws.page_setup.fitToHeight = 0
            ws.print_options.horizontalCentered = True
            ws.print_title_rows = f"{riga_header}:{riga_header}"
            ws.oddFooter.center.text = "Pagina &P di &N"

    writer_classi.close()

def genera_orario_docenti(calendario, docenti_civics_df, cartella_output):
    """Genera il file orario_docenti.xlsx: un foglio per ogni docente, con una griglia ore x giorni per ogni settimana"""

    colore_scuro = _COLORE_SCURO
    lato = Side(style='thin', color=_COLORE_BORDO)
    bordo = Border(left=lato, right=lato, top=lato, bottom=lato)
    settimana_fill = PatternFill(start_color=colore_scuro, end_color=colore_scuro, fill_type='solid')
    giorni_fill = PatternFill(start_color='D9E2F3', end_color='D9E2F3', fill_type='solid')
    ora_fill = PatternFill(start_color='EEF3FA', end_color='EEF3FA', fill_type='solid')
    lezione_fill = PatternFill(start_color='FFF2CC', end_color='FFF2CC', fill_type='solid')
    centrato_a_capo = Alignment(horizontal='center', vertical='center', wrap_text=True)

    writer_docenti = pd.ExcelWriter(os.path.join(cartella_output, 'orario_docenti.xlsx'), engine='openpyxl')

    giorni_tutti = ['LUN', 'MAR', 'MER', 'GIO', 'VEN', 'SAB']
    riga_iniziale = 4
    settimane_per_pagina = 3

    # Pre-raggruppamento delle voci per docente per ottimizzare la ricerca
    calendario_per_docente = defaultdict(list)
    for entry in calendario:
        calendario_per_docente[entry['DOCENTE_CIVICS']].append(entry)

    # Per ogni docente di civics creiamo un foglio con una griglia (ore x giorni) per ogni settimana
    for nome_docente in docenti_civics_df['DOCENTE']:
        docente_entries = calendario_per_docente.get(nome_docente, [])
        if not docente_entries:
            continue

        # Lezioni raggruppate per settimana (identificata dal lunedì, così l'ordine è cronologico)
        lezioni_per_settimana = defaultdict(dict)
        for entry in docente_entries:
            lunedi = entry['DATA'] - timedelta(days=entry['DATA'].weekday())
            lezioni_per_settimana[lunedi][(entry['GIORNO'], entry['ORA'])] = \
                f"{entry['CLASSE']}\n({entry['DOCENTE_SOSTITUITO']})"

        # Il sabato compare solo se il docente ha almeno una lezione quel giorno; le ore sono almeno 6
        giorni = [g for g in giorni_tutti if g != 'SAB' or any(e['GIORNO'] == 'SAB' for e in docente_entries)]
        n_ore = max(6, max(e['ORA'] for e in docente_entries))
        n_colonne = len(giorni) + 1

        # Costruzione delle righe: per ogni settimana titolo, intestazione dei giorni, ore e riga vuota
        righe = []
        tipi = []
        for numero, lunedi in enumerate(sorted(lezioni_per_settimana), start=1):
            righe.append([f"Settimana {numero}  ·  {_get_week_range(lunedi)}"] + [''] * len(giorni))
            tipi.append('settimana')
            righe.append(['Ora'] + [f"{_NOMI_GIORNI[g]}\n{(lunedi + timedelta(days=giorni_tutti.index(g))).strftime('%d/%m')}"
                                    for g in giorni])
            tipi.append('giorni')
            for ora in range(1, n_ore + 1):
                righe.append([ora] + [lezioni_per_settimana[lunedi].get((g, ora), '') for g in giorni])
                tipi.append('ora')
            righe.append([''] * n_colonne)
            tipi.append('vuota')

        df = pd.DataFrame(righe)
        df = _sanitize_for_excel(df)
        safe_sheet_name = _sanitize_sheet_name(nome_docente, default="Docente")
        df.to_excel(writer_docenti, sheet_name=safe_sheet_name, index=False, header=False, startrow=riga_iniziale - 1)

        ws = writer_docenti.sheets[safe_sheet_name]

        _scrivi_titolo_foglio(
            ws, f"Educazione civica – Docente {nome_docente}",
            f"{len(docente_entries)} ore  ·  {len({e['CLASSE'] for e in docente_entries})} classi  ·  "
            f"{len(lezioni_per_settimana)} settimane  ·  dal {min(e['DATA'] for e in docente_entries).strftime('%d/%m/%Y')} "
            f"al {max(e['DATA'] for e in docente_entries).strftime('%d/%m/%Y')}",
            n_colonne)

        # Formattazione per tipo di riga
        ws.column_dimensions['A'].width = 9
        for col in range(2, n_colonne + 1):
            ws.column_dimensions[get_column_letter(col)].width = 22
        for indice, tipo in enumerate(tipi):
            riga = riga_iniziale + indice
            if tipo == 'settimana':
                ws.merge_cells(start_row=riga, start_column=1, end_row=riga, end_column=n_colonne)
                cella = ws.cell(row=riga, column=1)
                cella.fill = settimana_fill
                cella.font = Font(bold=True, color='FFFFFF')
                cella.alignment = Alignment(horizontal='left', vertical='center', indent=1)
                ws.row_dimensions[riga].height = 22
            elif tipo == 'giorni':
                for col in range(1, n_colonne + 1):
                    cella = ws.cell(row=riga, column=col)
                    cella.fill = giorni_fill
                    cella.font = Font(bold=True)
                    cella.border = bordo
                    cella.alignment = centrato_a_capo
                ws.row_dimensions[riga].height = 32
            elif tipo == 'ora':
                for col in range(1, n_colonne + 1):
                    cella = ws.cell(row=riga, column=col)
                    cella.border = bordo
                    cella.alignment = centrato_a_capo
                    if col == 1:
                        cella.fill = ora_fill
                        cella.font = Font(bold=True)
                    elif cella.value:
                        cella.fill = lezione_fill
                ws.row_dimensions[riga].height = 32
            else:
                ws.row_dimensions[riga].height = 10

        # Vista e stampa: A4 in larghezza, 3 settimane per pagina (interruzione dopo la riga vuota)
        righe_per_settimana = n_ore + 3
        for numero_settimana in range(settimane_per_pagina, len(lezioni_per_settimana), settimane_per_pagina):
            ws.row_breaks.append(Break(id=riga_iniziale + numero_settimana * righe_per_settimana - 1))
        ws.sheet_view.showGridLines = False
        ws.sheet_properties.tabColor = colore_scuro
        ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
        ws.page_setup.orientation = 'portrait'
        ws.page_setup.paperSize = ws.PAPERSIZE_A4
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.print_options.horizontalCentered = True
        ws.oddFooter.center.text = "Pagina &P di &N"

    writer_docenti.close()

def genera_file_excel(calendario, classi_df, docenti_civics_df, cartella_output):
    # Riduce il percorso a un nome di cartella per evitare path traversal.
    cartella_output = _sanitize_output_path(cartella_output)

    # Questa funzione genera due file Excel:
    # 1. orario_classi.xlsx: un foglio per ogni classe, con le sostituzioni settimanali
    # 2. orario_docenti.xlsx: un foglio per ogni docente, con le ore settimanali su righe

    # Converte le date in datetime, se stringhe (ottimizzato con cache)
    date_cache = {}
    for entry in calendario:
        if isinstance(entry['DATA'], str):
            date_str = entry['DATA']
            if date_str not in date_cache:
                date_cache[date_str] = datetime.strptime(date_str, '%d/%m/%Y')
            entry['DATA'] = date_cache[date_str]

    genera_orario_classi(calendario, classi_df, cartella_output)
    genera_orario_docenti(calendario, docenti_civics_df, cartella_output)

